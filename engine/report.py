"""Outputs: console summary, CSV files and a formatted Excel workbook."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

PCT_HINTS = ("return", "rate", "mfe", "mae", "excess", "share", "win", "loss", "expectancy",
             "median", "worst", "best", "hit", "exit", "risk_pct", "breadth", "ret_", "avg_max")

NOTES = [
    ("What this is", "Point-in-time event study: rules are applied on each past scan date using only data "
                     "known that day; entries at the next session's open; outcomes measured afterwards."),
    ("Max upside (MFE)", "Highest high within the horizon vs entry - the 'upside potential' that was available. "
                         "Not a return you would have captured; see Trades for rule-based exits."),
    ("Excess vs universe", "Pick return minus the average return of every liquid stock on the same date. "
                           "This is the real test: does the screen beat picking any liquid stock?"),
    ("Control group", "control_naive_low = liquid stocks within 5% of their 52W low. Tests whether buying lows "
                      "blindly works; the reversal model must beat it to justify its rules."),
    ("Survivorship bias", "The universe is TODAY's index list, so companies that fell out or delisted are missing. "
                          "Results are optimistic - use a point-in-time constituent history before trusting them."),
    ("Not in v1", "Delivery %, fundamentals, results dates, pledges, news/catalysts, sector indices "
                  "(sector strength is derived from universe members). Their scoring weights were re-distributed."),
    ("Data", "Adjusted EOD prices from the configured source. Yahoo data can contain errors; spot-check "
             "against NSE for any stock you act on."),
    ("Overfitting", "Thresholds are starting assumptions. Tune only on the in-sample period and judge on "
                    "the out-of-sample period; do not tune on the latest few months."),
    ("Not advice", "Research tool output. Not a recommendation to buy or sell any security."),
]


def print_summary(res: dict, picks: pd.DataFrame | None = None) -> None:
    pd.set_option("display.width", 200, "display.max_columns", 30)
    s = res.get("summary", {})
    if "forward" in s:
        f = s["forward"][["model", "horizon_days", "signals", "avg_return", "hit_rate_positive",
                          "avg_max_upside_MFE", "share_reaching_+5pct", "avg_max_drawdown_MAE",
                          "universe_avg_return", "avg_excess_vs_universe"]]
        print("\n=== Forward outcomes of picks (entry next open) ===")
        print(f.to_string(index=False, float_format=lambda v: f"{v:.3f}"))
    if "trades" in s:
        print("\n=== Simulated trades (stop / target / time exit, after costs) ===")
        print(s["trades"].to_string(index=False, float_format=lambda v: f"{v:.3f}"))
    if "by_period" in s:
        print("\n=== By period (in-sample vs out-of-sample) ===")
        print(s["by_period"][["model", "period", "trades", "win_rate", "expectancy", "profit_factor"]]
              .to_string(index=False, float_format=lambda v: f"{v:.3f}"))
    if picks is not None and len(picks):
        print(f"\n=== Latest lists (as of {picks['as_of'].iloc[0]}, regime: {picks['regime'].iloc[0]}) ===")
        print(picks[["model", "rank", "symbol", "setup", "score", "last_close", "stop_ref", "risk_pct"]]
              .to_string(index=False))


def _format(path: Path) -> None:
    wb = load_workbook(path)
    head_fill = PatternFill("solid", start_color="1F3864")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.row == 1, color="FFFFFF" if c.row == 1 else "000000")
        for c in ws[1]:
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        headers = [str(c.value or "").lower() for c in ws[1]]
        for idx, h in enumerate(headers, start=1):
            col = ws.cell(row=1, column=idx).column_letter
            is_pct = any(k in h for k in PCT_HINTS) and not any(k in h for k in ("days", "trades", "signals"))
            width = 14
            if h in ("why", "risk_flags", "note"):
                width = 70
            elif h in ("detail",):
                width = 110
            ws.column_dimensions[col].width = width
            for c in ws[col][1:]:
                if isinstance(c.value, float):
                    c.number_format = "0.0%" if is_pct else "0.00"
                if h in ("why", "risk_flags", "detail"):
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    wb.save(path)


def write_outputs(out_dir: str | Path, res: dict | None, picks: pd.DataFrame | None, cfg: dict,
                  tag: str) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    xlsx = out / f"momentum_engine_{tag}.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        notes = pd.DataFrame(NOTES, columns=["topic", "detail"])
        if res:
            notes.loc[len(notes)] = ["Weights used", str(res.get("weights_used"))]
            notes.loc[len(notes)] = ["Scan dates", f"{len(res['scan_dates'])} scans, "
                                     f"{res['scan_dates'][0].date()} to {res['scan_dates'][-1].date()}"]
        notes.to_excel(xw, sheet_name="Read_me", index=False)
        if picks is not None and len(picks):
            for model in picks["model"].unique():
                picks[picks["model"] == model].to_excel(xw, sheet_name=f"Latest_{model}"[:31], index=False)
        if res:
            for name, df in res["summary"].items():
                df.to_excel(xw, sheet_name=f"Summary_{name}"[:31], index=False)
            if len(res["trades"]):
                res["trades"].to_excel(xw, sheet_name="Trades", index=False)
            res["regime"].dropna(subset=["nifty"]).iloc[::5].rename_axis("date").to_excel(xw, sheet_name="Regime_weekly")
    _format(xlsx)
    if res:
        res["signals"].to_csv(out / f"signals_{tag}.csv", index=False)
        res["trades"].to_csv(out / f"trades_{tag}.csv", index=False)
    if picks is not None and len(picks):
        picks.to_csv(out / f"picks_{picks['as_of'].iloc[0]}.csv", index=False)
    return xlsx
